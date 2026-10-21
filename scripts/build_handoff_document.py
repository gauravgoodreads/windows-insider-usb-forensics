from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_HANDOFF.md"
OUTPUT = ROOT / "docs" / "Forensic_Project_Handoff_and_Demonstration_Guide.docx"
IMAGE_DIR = ROOT / "docs" / "images"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / name), size)


def rounded_box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], text: str, fill: str) -> None:
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline="#17365D", width=4)
    x1, y1, x2, y2 = xy
    lines = text.split("\n")
    f = font(31, True)
    total_h = len(lines) * 39
    y = (y1 + y2 - total_h) / 2
    for line in lines:
        box = draw.textbbox((0, 0), line, font=f)
        draw.text(((x1 + x2 - (box[2] - box[0])) / 2, y), line, font=f, fill="#111111")
        y += 39


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line((start, end), fill="#17365D", width=8)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        sign = 1 if ex > sx else -1
        points = [(ex, ey), (ex - sign * 24, ey - 15), (ex - sign * 24, ey + 15)]
    else:
        sign = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 15, ey - sign * 24), (ex + 15, ey - sign * 24)]
    draw.polygon(points, fill="#17365D")


def make_diagrams() -> list[Path]:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []

    image = Image.new("RGB", (1800, 720), "white")
    d = ImageDraw.Draw(image)
    labels = [
        "Scenario data\nor real PC",
        "Read only collection\nand hashes",
        "Read only\nexports",
        "Normalize and\ncorrelate",
        "Hypotheses and\nSQLite",
        "Dashboard and\nreports",
    ]
    fills = ["#D6EAF8", "#D5F5E3", "#FCF3CF", "#E8DAEF", "#FADBD8", "#D6DBDF"]
    x_positions = [50, 345, 640, 935, 1230, 1525]
    for i, (label, fill, x) in enumerate(zip(labels, fills, x_positions)):
        rounded_box(d, (x, 250, x + 225, 445), label, fill)
        if i < len(labels) - 1:
            arrow(d, (x + 225, 348), (x_positions[i + 1] - 15, 348))
    d.text((50, 70), "Complete Forensic Workflow", font=font(48, True), fill="#000000")
    d.text((50, 135), "Controlled activity becomes explainable evidence through separated trust stages", font=font(29), fill="#333333")
    path = IMAGE_DIR / "handoff_system_flow.png"
    image.save(path)
    outputs.append(path)

    image = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(image)
    rounded_box(d, (560, 205, 1040, 340), "Preserved source\nand verified copies", "#D5F5E3")
    rounded_box(d, (80, 420, 455, 585), "Scenarios and collection\nknown safe actions", "#D6EAF8")
    rounded_box(d, (610, 420, 990, 585), "Analysis\nraw to normalized", "#E8DAEF")
    rounded_box(d, (1145, 420, 1520, 585), "Presentation\nreasoning and limits", "#FADBD8")
    rounded_box(d, (560, 710, 1040, 845), "Traceability\nartifact IDs and raw payload", "#FCF3CF")
    arrow(d, (455, 502), (595, 502))
    arrow(d, (990, 502), (1130, 502))
    arrow(d, (800, 340), (800, 405))
    arrow(d, (800, 585), (800, 695))
    d.text((80, 70), "Trust Boundaries", font=font(48, True), fill="#000000")
    d.text((80, 135), "Simulation metadata is context, not evidence; hashes and provenance connect the remaining stages", font=font(27), fill="#333333")
    path = IMAGE_DIR / "handoff_trust_boundaries.png"
    image.save(path)
    outputs.append(path)

    image = Image.new("RGB", (1700, 920), "white")
    d = ImageDraw.Draw(image)
    inputs = [
        (80, 190, "USB device\nactivity"),
        (80, 360, "Sensitive file\nactivity"),
        (80, 530, "Destination side\nevidence"),
        (80, 700, "User archive and\ndeletion traces"),
    ]
    for x, y, label in inputs:
        rounded_box(d, (x, y, x + 390, y + 125), label, "#EAF2F8")
        arrow(d, (x + 390, y + 62), (660, 455))
    rounded_box(d, (675, 330, 1075, 580), "Configured time\nand identity rules", "#FCF3CF")
    arrow(d, (1075, 455), (1215, 455))
    rounded_box(d, (1230, 330, 1630, 580), "Explainable\ncorrelations", "#E8DAEF")
    arrow(d, (1430, 580), (1430, 705))
    rounded_box(d, (1200, 720, 1660, 865), "H1 to H4 with evidence\nconfidence and limitations", "#D5F5E3")
    d.text((80, 55), "Corroboration Model", font=font(48, True), fill="#000000")
    d.text((80, 120), "No single trace proves transfer; independent observations must agree", font=font(29), fill="#333333")
    path = IMAGE_DIR / "handoff_corroboration_model.png"
    image.save(path)
    outputs.append(path)
    return outputs


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), GRID)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = BLACK

    for name, size in (("Heading 1", 19), ("Heading 2", 15), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])
        for run in section.footer.paragraphs[0].runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 90, 90)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    p.add_run("Complete Project Handoff and Demonstration Guide")
    p_pr = p._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    borders.append(bottom)
    p_pr.append(borders)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("USB Insider Forensics: Detecting Data Exfiltration via Removable Media on Windows")
    r.bold = True
    r.font.size = Pt(15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.add_run("Knowledge transfer, continuation procedure, architecture, verification, and examiner demonstration").font.size = Pt(12)

    details = [
        "Case: INSIDER-USB-001",
        "Repository: gauravgoodreads/windows-insider-usb-forensics",
        r"Local project: C:\Users\Gary\Desktop\CF proj",
        "Prepared: 21 October 2026",
    ]
    for value in details:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(value).font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    r = p.add_run("Application status: validated and demonstration ready\nEvidence: controlled scenario dataset plus a real run of Experiments A and B (no virtual machine needed)")
    r.bold = True
    r.font.size = Pt(11)
    doc.add_page_break()


def add_toc(doc: Document, headings: list[str]) -> None:
    doc.add_heading("Document Guide", level=1)
    p = doc.add_paragraph("Use this guide to explain the project, collect real evidence, verify the implementation, and demonstrate the results safely.")
    p.paragraph_format.space_after = Pt(10)
    for title in headings:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(title)
    doc.add_page_break()


def format_run(run, bold: bool | None = None, italic: bool | None = None, code: bool = False) -> None:
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if code:
        run.font.name = "Cascadia Mono"
        run.font.size = Pt(9)


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            format_run(paragraph.add_run(token[1:-1]), code=True)
        else:
            format_run(paragraph.add_run(token[2:-2]), bold=True)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GRAY)
    p_pr.append(shd)
    run = p.add_run("\n".join(lines))
    format_run(run, code=True)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = True
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        prevent_row_split(table.rows[ridx])
        for cidx in range(cols):
            cell = table.cell(ridx, cidx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = row[cidx].strip() if cidx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
            for run in p.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8.5)
                if ridx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if ridx == 0:
                set_cell_shading(cell, NAVY)
            elif ridx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table_line(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def build_document(diagrams: list[Path]) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]
    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_toc(doc, headings)

    i = 1  # Skip Markdown title because the cover already contains it.
    diagram_index = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("## "):
            heading_text = line[3:].strip()
            if heading_text == "Full repository file inventory":
                doc.add_page_break()
            doc.add_heading(heading_text, level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue
        if line.startswith("```mermaid"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                i += 1
            if diagram_index < len(diagrams):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(diagrams[diagram_index]), width=Inches(7.0))
                cap = doc.add_paragraph(f"Figure {diagram_index + 1}  {['Complete forensic workflow', 'Trust boundaries', 'Corroboration model'][diagram_index]}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True
                cap.runs[0].font.size = Pt(9)
                diagram_index += 1
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1]):
            table_rows = [parse_table_line(line)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_rows.append(parse_table_line(lines[i]))
                i += 1
            add_table(doc, table_rows)
            continue
        if re.match(r"^- ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            number, content = line.split(". ", 1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.add_run(f"{number}. ")
            add_inline(p, content.strip())
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            add_inline(p, line[2:].strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        paragraph_lines = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{2,4} |```|\||- |\d+\. |> )", lines[i]):
            paragraph_lines.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))

    doc.core_properties.title = "Complete Project Handoff and Demonstration Guide"
    doc.core_properties.subject = "Windows USB cyber forensics project knowledge transfer"
    doc.core_properties.author = "Student Investigator"
    doc.core_properties.keywords = "cyber forensics, USB, Windows, evidence, handoff, demonstration"
    doc.save(OUTPUT)


def main() -> int:
    diagrams = make_diagrams()
    build_document(diagrams)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
